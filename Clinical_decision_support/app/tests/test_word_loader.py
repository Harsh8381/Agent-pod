from docx import Document

from app.rag.chunker import chunk_word_blocks
from app.rag.word_loader import load_word_blocks


def test_load_word_blocks_preserves_headings_and_tables(tmp_path):
    path = tmp_path / "guidelines.docx"
    document = Document()
    document.add_heading("Hypertension", level=1)
    document.add_paragraph("Measure blood pressure twice.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Target"
    table.rows[0].cells[1].text = "< 140/90"
    document.save(path)

    blocks = load_word_blocks(path)

    assert blocks[0]["heading"] == "Hypertension"
    assert blocks[1]["heading"] == "Hypertension"
    assert "Target | < 140/90" in blocks[2]["text"]


def test_chunk_word_blocks_adds_source_and_block_metadata():
    chunks = chunk_word_blocks(
        [{"text": "Clinical guidance.", "block": 7, "heading": "Asthma"}],
        "guidelines.docx",
    )

    assert chunks[0]["metadata"] == {
        "source": "guidelines.docx",
        "block": "7",
        "chunk": 1,
        "heading": "Asthma",
    }