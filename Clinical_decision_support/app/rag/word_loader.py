from pathlib import Path

from docx import Document


def load_word_blocks(docx_path):
    docx_path = Path(docx_path)
    if not docx_path.exists():
        raise FileNotFoundError(f"Word document not found: {docx_path}")

    document = Document(str(docx_path))
    blocks = []
    current_heading = ""

    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            current_heading = text
        blocks.append({
            "text": text,
            "block": paragraph_number,
            "heading": current_heading,
        })

    for table_number, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append({
                "text": "\n".join(rows),
                "block": f"table-{table_number}",
                "heading": current_heading,
            })
    return blocks